"""
Tests para el modelo de exportación a Excel.
"""

import pytest
import os
from datetime import datetime
from openpyxl import load_workbook
from models.export_excel_model import ExcelExportModel


@pytest.fixture
def export_model():
    """Fixture que proporciona un modelo de exportación Excel."""
    return ExcelExportModel()


@pytest.fixture
def custom_export_model():
    """Fixture con configuración personalizada."""
    return ExcelExportModel(
        title='Test Export',
        author='Test Author',
        include_charts=False
    )


@pytest.fixture
def sample_history():
    """Fixture con historial de ejemplo."""
    return [
        {
            'timestamp': '2024-01-15T10:30:00',
            'expression': '2 + 2',
            'result': '4',
            'mode': 'basic',
            'notes': 'Suma simple'
        },
        {
            'timestamp': '2024-01-15T10:31:00',
            'expression': 'sin(pi/2)',
            'result': '1.0',
            'mode': 'scientific',
            'notes': ''
        },
        {
            'timestamp': '2024-01-15T10:32:00',
            'expression': '10 / 0',
            'result': '',
            'mode': 'basic',
            'error': 'División por cero',
            'notes': 'Error esperado'
        }
    ]


@pytest.fixture
def large_history():
    """Fixture con historial grande."""
    entries = []
    modes = ['basic', 'scientific', 'extended']
    for i in range(100):
        entries.append({
            'timestamp': f'2024-01-15T{10 + i//60:02d}:{i%60:02d}:00',
            'expression': f'{i} + {i+1}',
            'result': str(2*i + 1),
            'mode': modes[i % 3],
            'notes': f'Cálculo {i+1}'
        })
    return entries


class TestExcelExportModelInit:
    """Tests para inicialización del modelo."""
    
    def test_default_initialization(self):
        """Test inicialización con valores por defecto."""
        model = ExcelExportModel()
        assert model.title == 'Historial de Cálculos'
        assert model.author == 'Calculadora Científica'
        assert model.include_charts is True
    
    def test_custom_initialization(self):
        """Test inicialización con valores personalizados."""
        model = ExcelExportModel(
            title='Mi Historial',
            author='Usuario Test',
            include_charts=False
        )
        assert model.title == 'Mi Historial'
        assert model.author == 'Usuario Test'
        assert model.include_charts is False
    
    def test_str_representation(self):
        """Test representación en string."""
        model = ExcelExportModel()
        result = str(model)
        assert 'ExcelExportModel' in result
        assert 'Historial de Cálculos' in result
        assert 'include_charts=True' in result
    
    def test_repr_representation(self):
        """Test representación detallada."""
        model = ExcelExportModel(title='Test', author='Author')
        result = repr(model)
        assert 'ExcelExportModel' in result
        assert "'Test'" in result
        assert "'Author'" in result


class TestExportHistory:
    """Tests para exportación de historial."""
    
    def test_export_basic_history(self, export_model, sample_history, tmp_path):
        """Test exportación básica de historial."""
        filename = str(tmp_path / "test_history.xlsx")
        result = export_model.export_history(sample_history, filename)
        
        assert result == filename
        assert os.path.exists(filename)
        
        # Verificar contenido
        wb = load_workbook(filename)
        assert 'Historial' in wb.sheetnames
        assert 'Estadísticas' in wb.sheetnames
    
    def test_export_adds_xlsx_extension(self, export_model, sample_history, tmp_path):
        """Test que se agrega extensión .xlsx automáticamente."""
        filename = str(tmp_path / "test_history")
        result = export_model.export_history(sample_history, filename)
        
        assert result.endswith('.xlsx')
        assert os.path.exists(result)
    
    def test_export_without_statistics(self, export_model, sample_history, tmp_path):
        """Test exportación sin hoja de estadísticas."""
        filename = str(tmp_path / "test_no_stats.xlsx")
        export_model.export_history(sample_history, filename, include_statistics=False)
        
        wb = load_workbook(filename)
        assert 'Historial' in wb.sheetnames
        assert 'Estadísticas' not in wb.sheetnames
    
    def test_export_empty_history_raises_error(self, export_model, tmp_path):
        """Test que exportar historial vacío lanza error."""
        filename = str(tmp_path / "empty.xlsx")
        with pytest.raises(ValueError, match="No hay entradas para exportar"):
            export_model.export_history([], filename)
    
    def test_export_history_sheet_structure(self, export_model, sample_history, tmp_path):
        """Test estructura de la hoja de historial."""
        filename = str(tmp_path / "test_structure.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Historial']
        
        # Verificar encabezados
        assert ws['A1'].value == '#'
        assert ws['B1'].value == 'Fecha/Hora'
        assert ws['C1'].value == 'Expresión'
        assert ws['D1'].value == 'Resultado'
        assert ws['E1'].value == 'Modo'
        assert ws['F1'].value == 'Notas'
        
        # Verificar datos
        assert ws['A2'].value == 1
        assert ws['C2'].value == '2 + 2'
        assert ws['D2'].value == '4'
        assert ws['E2'].value == 'Basic'
    
    def test_export_handles_errors(self, export_model, sample_history, tmp_path):
        """Test que los errores se manejan correctamente."""
        filename = str(tmp_path / "test_errors.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Historial']
        
        # Verificar entrada con error (fila 4)
        error_cell = ws['D4'].value
        assert 'Error' in error_cell
        assert 'División por cero' in error_cell
    
    def test_export_large_history(self, export_model, large_history, tmp_path):
        """Test exportación de historial grande."""
        filename = str(tmp_path / "test_large.xlsx")
        result = export_model.export_history(large_history, filename)
        
        assert os.path.exists(result)
        
        wb = load_workbook(filename)
        ws = wb['Historial']
        
        # Verificar que todas las entradas están presentes
        # +1 por el encabezado
        assert ws.max_row == len(large_history) + 1
    
    def test_export_document_properties(self, export_model, sample_history, tmp_path):
        """Test propiedades del documento."""
        filename = str(tmp_path / "test_props.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        assert wb.properties.title == 'Historial de Cálculos'
        assert wb.properties.creator == 'Calculadora Científica'
        assert wb.properties.created is not None


class TestStatisticsSheet:
    """Tests para la hoja de estadísticas."""
    
    def test_statistics_sheet_created(self, export_model, sample_history, tmp_path):
        """Test que se crea la hoja de estadísticas."""
        filename = str(tmp_path / "test_stats.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        assert 'Estadísticas' in wb.sheetnames
    
    def test_statistics_general_info(self, export_model, sample_history, tmp_path):
        """Test información general en estadísticas."""
        filename = str(tmp_path / "test_stats_info.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Estadísticas']
        
        # Verificar título
        assert 'Estadísticas' in ws['A1'].value
        
        # Verificar total de cálculos
        assert ws['A4'].value == 'Total de cálculos:'
        assert ws['B4'].value == 3
    
    def test_statistics_success_rate(self, export_model, sample_history, tmp_path):
        """Test cálculo de tasa de éxito."""
        filename = str(tmp_path / "test_success.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Estadísticas']
        
        # Buscar tasa de éxito
        found = False
        for row in ws.iter_rows(min_row=7, max_row=12):
            if row[0].value == 'Tasa de éxito':
                assert '66.7%' in str(row[1].value) or '67%' in str(row[1].value)
                found = True
                break
        assert found, "No se encontró la tasa de éxito"
    
    def test_statistics_mode_distribution(self, export_model, sample_history, tmp_path):
        """Test distribución por modo."""
        filename = str(tmp_path / "test_modes.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Estadísticas']
        
        # Verificar tabla de modos
        assert 'Uso por Modo' in ws['A13'].value
        assert ws['A14'].value == 'Modo'
        assert ws['B14'].value == 'Cantidad'
        assert ws['C14'].value == 'Porcentaje'
    
    def test_statistics_with_charts(self, export_model, sample_history, tmp_path):
        """Test que se incluyen gráficos."""
        filename = str(tmp_path / "test_charts.xlsx")
        export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Estadísticas']
        
        # Verificar que hay gráficos
        assert len(ws._charts) > 0
    
    def test_statistics_without_charts(self, custom_export_model, sample_history, tmp_path):
        """Test estadísticas sin gráficos."""
        filename = str(tmp_path / "test_no_charts.xlsx")
        custom_export_model.export_history(sample_history, filename)
        
        wb = load_workbook(filename)
        ws = wb['Estadísticas']
        
        # Verificar que no hay gráficos
        assert len(ws._charts) == 0


class TestExportSingleCalculation:
    """Tests para exportación de cálculo individual."""
    
    def test_export_single_basic(self, export_model, tmp_path):
        """Test exportación de cálculo individual básico."""
        filename = str(tmp_path / "single.xlsx")
        result = export_model.export_single_calculation(
            expression='5 + 3',
            result='8',
            filename=filename
        )
        
        assert result == filename
        assert os.path.exists(filename)
    
    def test_export_single_adds_extension(self, export_model, tmp_path):
        """Test que se agrega extensión .xlsx."""
        filename = str(tmp_path / "single")
        result = export_model.export_single_calculation(
            expression='2 * 3',
            result='6',
            filename=filename
        )
        
        assert result.endswith('.xlsx')
        assert os.path.exists(result)
    
    def test_export_single_with_mode(self, export_model, tmp_path):
        """Test exportación con modo específico."""
        filename = str(tmp_path / "single_mode.xlsx")
        export_model.export_single_calculation(
            expression='sin(0)',
            result='0',
            filename=filename,
            mode='scientific'
        )
        
        wb = load_workbook(filename)
        ws = wb.active
        
        # Buscar el modo
        found = False
        for row in ws.iter_rows():
            if row[0].value == 'Modo:':
                assert row[1].value == 'Scientific'
                found = True
                break
        assert found
    
    def test_export_single_with_notes(self, export_model, tmp_path):
        """Test exportación con notas."""
        filename = str(tmp_path / "single_notes.xlsx")
        notes = "Cálculo de prueba"
        export_model.export_single_calculation(
            expression='10 + 5',
            result='15',
            filename=filename,
            notes=notes
        )
        
        wb = load_workbook(filename)
        ws = wb.active
        
        # Buscar las notas
        found = False
        for row in ws.iter_rows():
            if row[0].value == 'Notas:':
                assert row[1].value == notes
                found = True
                break
        assert found
    
    def test_export_single_structure(self, export_model, tmp_path):
        """Test estructura del archivo de cálculo individual."""
        filename = str(tmp_path / "single_struct.xlsx")
        export_model.export_single_calculation(
            expression='7 * 8',
            result='56',
            filename=filename
        )
        
        wb = load_workbook(filename)
        ws = wb.active
        
        # Verificar título
        assert 'Cálculo Individual' in ws['A1'].value
        
        # Verificar que tiene fecha
        found_date = False
        for row in ws.iter_rows():
            if row[0].value == 'Fecha:':
                assert row[1].value is not None
                found_date = True
                break
        assert found_date
    
    def test_export_single_document_properties(self, export_model, tmp_path):
        """Test propiedades del documento individual."""
        filename = str(tmp_path / "single_props.xlsx")
        export_model.export_single_calculation(
            expression='1 + 1',
            result='2',
            filename=filename
        )
        
        wb = load_workbook(filename)
        assert wb.properties.title == 'Cálculo Individual'
        assert wb.properties.creator == 'Calculadora Científica'


class TestEdgeCases:
    """Tests para casos extremos."""
    
    def test_export_with_special_characters(self, export_model, tmp_path):
        """Test exportación con caracteres especiales."""
        history = [{
            'timestamp': '2024-01-15T10:30:00',
            'expression': 'π * 2',
            'result': '6.283185307179586',
            'mode': 'scientific',
            'notes': 'Usando π'
        }]
        
        filename = str(tmp_path / "special_chars.xlsx")
        export_model.export_history(history, filename)
        
        assert os.path.exists(filename)
        wb = load_workbook(filename)
        ws = wb['Historial']
        assert 'π' in ws['C2'].value
    
    def test_export_with_long_expressions(self, export_model, tmp_path):
        """Test exportación con expresiones largas."""
        long_expr = ' + '.join([str(i) for i in range(50)])
        history = [{
            'timestamp': '2024-01-15T10:30:00',
            'expression': long_expr,
            'result': str(sum(range(50))),
            'mode': 'basic',
            'notes': ''
        }]
        
        filename = str(tmp_path / "long_expr.xlsx")
        export_model.export_history(history, filename)
        
        assert os.path.exists(filename)
    
    def test_export_with_invalid_timestamp(self, export_model, tmp_path):
        """Test exportación con timestamp inválido."""
        history = [{
            'timestamp': 'invalid-timestamp',
            'expression': '1 + 1',
            'result': '2',
            'mode': 'basic',
            'notes': ''
        }]
        
        filename = str(tmp_path / "invalid_ts.xlsx")
        export_model.export_history(history, filename)
        
        assert os.path.exists(filename)
    
    def test_export_with_missing_fields(self, export_model, tmp_path):
        """Test exportación con campos faltantes."""
        history = [{
            'expression': '2 + 2',
            'result': '4'
        }]
        
        filename = str(tmp_path / "missing_fields.xlsx")
        export_model.export_history(history, filename)
        
        assert os.path.exists(filename)
    
    def test_export_with_empty_strings(self, export_model, tmp_path):
        """Test exportación con strings vacíos."""
        history = [{
            'timestamp': '',
            'expression': '',
            'result': '',
            'mode': '',
            'notes': ''
        }]
        
        filename = str(tmp_path / "empty_strings.xlsx")
        export_model.export_history(history, filename)
        
        assert os.path.exists(filename)


class TestIntegration:
    """Tests de integración."""
    
    def test_full_workflow(self, export_model, sample_history, tmp_path):
        """Test flujo completo de exportación."""
        # Exportar historial
        history_file = str(tmp_path / "full_history.xlsx")
        export_model.export_history(sample_history, history_file)
        
        # Exportar cálculo individual
        single_file = str(tmp_path / "full_single.xlsx")
        export_model.export_single_calculation(
            expression='test',
            result='result',
            filename=single_file
        )
        
        # Verificar ambos archivos
        assert os.path.exists(history_file)
        assert os.path.exists(single_file)
        
        # Verificar contenido
        wb1 = load_workbook(history_file)
        assert len(wb1.sheetnames) == 2
        
        wb2 = load_workbook(single_file)
        assert len(wb2.sheetnames) == 1
    
    def test_multiple_exports_same_model(self, export_model, sample_history, tmp_path):
        """Test múltiples exportaciones con el mismo modelo."""
        files = []
        for i in range(3):
            filename = str(tmp_path / f"export_{i}.xlsx")
            export_model.export_history(sample_history, filename)
            files.append(filename)
        
        # Verificar todos los archivos
        for f in files:
            assert os.path.exists(f)
    
    def test_custom_model_workflow(self, tmp_path):
        """Test flujo con modelo personalizado."""
        model = ExcelExportModel(
            title='Custom Export',
            author='Test User',
            include_charts=False
        )
        
        history = [{
            'timestamp': '2024-01-15T10:30:00',
            'expression': 'test',
            'result': 'result',
            'mode': 'basic',
            'notes': ''
        }]
        
        filename = str(tmp_path / "custom.xlsx")
        model.export_history(history, filename)
        
        wb = load_workbook(filename)
        assert wb.properties.title == 'Custom Export'
        assert wb.properties.creator == 'Test User'
        
        # Verificar que no hay gráficos
        ws = wb['Estadísticas']
        assert len(ws._charts) == 0

# Made with Bob
