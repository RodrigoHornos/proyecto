"""
Modelo para exportación de historial a Excel.

Este módulo proporciona funcionalidades para exportar el historial
de cálculos a hojas de cálculo Excel con formato profesional.
"""

from typing import List, Dict, Optional
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference


class ExcelExportModel:
    """
    Modelo para exportar historial de cálculos a Excel.
    
    Attributes:
        title: Título del documento
        author: Autor del documento
        include_charts: Incluir gráficos
    """
    
    def __init__(
        self,
        title: str = 'Historial de Cálculos',
        author: str = 'Calculadora Científica',
        include_charts: bool = True
    ):
        """
        Inicializa el modelo de exportación Excel.
        
        Args:
            title: Título del documento
            author: Autor del documento
            include_charts: Incluir gráficos en la exportación
        """
        self.title = title
        self.author = author
        self.include_charts = include_charts
    
    def export_history(
        self,
        history_entries: List[Dict],
        filename: str,
        include_statistics: bool = True
    ) -> str:
        """
        Exporta el historial de cálculos a un archivo Excel.
        
        Args:
            history_entries: Lista de entradas del historial
            filename: Nombre del archivo Excel de salida
            include_statistics: Incluir hoja de estadísticas
            
        Returns:
            Ruta del archivo generado
            
        Raises:
            ValueError: Si no hay entradas para exportar
        """
        if not history_entries:
            raise ValueError("No hay entradas para exportar")
        
        # Asegurar extensión .xlsx
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        # Crear libro de trabajo
        wb = Workbook()
        
        # Eliminar hoja por defecto
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        # Crear hoja de historial
        self._create_history_sheet(wb, history_entries)
        
        # Crear hoja de estadísticas si se solicita
        if include_statistics:
            self._create_statistics_sheet(wb, history_entries)
        
        # Configurar propiedades del documento
        wb.properties.title = self.title
        wb.properties.creator = self.author
        wb.properties.created = datetime.now()
        
        # Guardar archivo
        wb.save(filename)
        
        return filename
    
    def _create_history_sheet(self, wb: Workbook, entries: List[Dict]) -> None:
        """
        Crea la hoja de historial.
        
        Args:
            wb: Libro de trabajo
            entries: Lista de entradas del historial
        """
        ws = wb.create_sheet("Historial", 0)
        
        # Estilos
        header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        cell_alignment = Alignment(horizontal='left', vertical='center')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Encabezados
        headers = ['#', 'Fecha/Hora', 'Expresión', 'Resultado', 'Modo', 'Notas']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # Datos
        for i, entry in enumerate(entries, 2):
            # Número
            cell = ws.cell(row=i, column=1, value=i-1)
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
            
            # Fecha/Hora
            timestamp = entry.get('timestamp', '')
            if timestamp:
                try:
                    dt = datetime.fromisoformat(timestamp)
                    timestamp = dt.strftime('%d/%m/%Y %H:%M')
                except:
                    pass
            cell = ws.cell(row=i, column=2, value=timestamp)
            cell.alignment = cell_alignment
            cell.border = border
            
            # Expresión
            expression = entry.get('expression', '')
            cell = ws.cell(row=i, column=3, value=expression)
            cell.alignment = cell_alignment
            cell.border = border
            
            # Resultado
            result = entry.get('result', '')
            if entry.get('error'):
                result = f"Error: {entry['error']}"
                cell = ws.cell(row=i, column=4, value=result)
                cell.font = Font(color='FF0000')
            else:
                cell = ws.cell(row=i, column=4, value=result)
            cell.alignment = Alignment(horizontal='right')
            cell.border = border
            
            # Modo
            mode = entry.get('mode', 'basic')
            cell = ws.cell(row=i, column=5, value=mode.capitalize())
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
            
            # Notas
            notes = entry.get('notes', '')
            cell = ws.cell(row=i, column=6, value=notes)
            cell.alignment = cell_alignment
            cell.border = border
            
            # Alternar colores de fila
            if i % 2 == 0:
                fill = PatternFill(start_color='ECF0F1', end_color='ECF0F1', fill_type='solid')
                for col in range(1, 7):
                    ws.cell(row=i, column=col).fill = fill
        
        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 12
        ws.column_dimensions['F'].width = 25
        
        # Congelar primera fila
        ws.freeze_panes = 'A2'
    
    def _create_statistics_sheet(self, wb: Workbook, entries: List[Dict]) -> None:
        """
        Crea la hoja de estadísticas.
        
        Args:
            wb: Libro de trabajo
            entries: Lista de entradas del historial
        """
        ws = wb.create_sheet("Estadísticas", 1)
        
        # Estilos
        title_font = Font(name='Arial', size=16, bold=True, color='2C3E50')
        header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='3498DB', end_color='3498DB', fill_type='solid')
        
        # Título
        ws['A1'] = 'Estadísticas del Historial'
        ws['A1'].font = title_font
        ws.merge_cells('A1:D1')
        
        # Información general
        ws['A3'] = 'Generado:'
        ws['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        ws['A4'] = 'Total de cálculos:'
        ws['B4'] = len(entries)
        
        # Calcular estadísticas
        total = len(entries)
        errors = sum(1 for e in entries if e.get('error'))
        success = total - errors
        
        # Estadísticas por modo
        modes = {}
        for entry in entries:
            mode = entry.get('mode', 'basic')
            modes[mode] = modes.get(mode, 0) + 1
        
        # Tabla de estadísticas generales
        ws['A6'] = 'Estadísticas Generales'
        ws['A6'].font = Font(name='Arial', size=14, bold=True)
        ws.merge_cells('A6:B6')
        
        ws['A7'] = 'Métrica'
        ws['B7'] = 'Valor'
        for cell in ['A7', 'B7']:
            ws[cell].font = header_font
            ws[cell].fill = header_fill
            ws[cell].alignment = Alignment(horizontal='center')
        
        stats_data = [
            ('Total de cálculos', total),
            ('Cálculos exitosos', success),
            ('Errores', errors),
            ('Tasa de éxito', f"{(success/total*100):.1f}%" if total > 0 else "0%")
        ]
        
        for i, (metric, value) in enumerate(stats_data, 8):
            ws[f'A{i}'] = metric
            ws[f'B{i}'] = value
            ws[f'B{i}'].alignment = Alignment(horizontal='center')
        
        # Tabla de uso por modo
        ws['A13'] = 'Uso por Modo'
        ws['A13'].font = Font(name='Arial', size=14, bold=True)
        ws.merge_cells('A13:C13')
        
        ws['A14'] = 'Modo'
        ws['B14'] = 'Cantidad'
        ws['C14'] = 'Porcentaje'
        for cell in ['A14', 'B14', 'C14']:
            ws[cell].font = header_font
            ws[cell].fill = header_fill
            ws[cell].alignment = Alignment(horizontal='center')
        
        row = 15
        for mode, count in sorted(modes.items(), key=lambda x: x[1], reverse=True):
            percentage = (count / total * 100) if total > 0 else 0
            ws[f'A{row}'] = mode.capitalize()
            ws[f'B{row}'] = count
            ws[f'B{row}'].alignment = Alignment(horizontal='center')
            ws[f'C{row}'] = f"{percentage:.1f}%"
            ws[f'C{row}'].alignment = Alignment(horizontal='center')
            row += 1
        
        # Ajustar anchos
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        
        # Agregar gráficos si está habilitado
        if self.include_charts and len(modes) > 0:
            self._add_charts(ws, modes, row)
    
    def _add_charts(self, ws, modes: Dict, start_row: int) -> None:
        """
        Agrega gráficos a la hoja de estadísticas.
        
        Args:
            ws: Hoja de trabajo
            modes: Diccionario de modos y cantidades
            start_row: Fila donde termina la tabla de modos
        """
        # Gráfico de barras
        chart = BarChart()
        chart.title = "Cálculos por Modo"
        chart.x_axis.title = "Modo"
        chart.y_axis.title = "Cantidad"
        
        data = Reference(ws, min_col=2, min_row=14, max_row=start_row-1)
        cats = Reference(ws, min_col=1, min_row=15, max_row=start_row-1)
        
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        
        ws.add_chart(chart, f"E6")
        
        # Gráfico de pastel
        pie = PieChart()
        pie.title = "Distribución por Modo"
        
        data = Reference(ws, min_col=2, min_row=14, max_row=start_row-1)
        cats = Reference(ws, min_col=1, min_row=15, max_row=start_row-1)
        
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(cats)
        
        ws.add_chart(pie, f"E20")
    
    def export_single_calculation(
        self,
        expression: str,
        result: str,
        filename: str,
        mode: str = 'basic',
        notes: Optional[str] = None
    ) -> str:
        """
        Exporta un solo cálculo a Excel.
        
        Args:
            expression: Expresión matemática
            result: Resultado del cálculo
            filename: Nombre del archivo Excel
            mode: Modo de cálculo
            notes: Notas adicionales
            
        Returns:
            Ruta del archivo generado
        """
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Cálculo"
        
        # Título
        ws['A1'] = 'Cálculo Individual'
        ws['A1'].font = Font(name='Arial', size=16, bold=True)
        ws.merge_cells('A1:B1')
        
        # Información
        ws['A3'] = 'Fecha:'
        ws['B3'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
        
        ws['A4'] = 'Modo:'
        ws['B4'] = mode.capitalize()
        
        ws['A6'] = 'Expresión:'
        ws['B6'] = expression
        ws['B6'].font = Font(name='Courier', size=12, bold=True, color='2980B9')
        
        ws['A8'] = 'Resultado:'
        ws['B8'] = result
        ws['B8'].font = Font(name='Courier', size=12, bold=True, color='27AE60')
        
        if notes:
            ws['A10'] = 'Notas:'
            ws['B10'] = notes
        
        # Ajustar anchos
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 40
        
        # Configurar propiedades
        wb.properties.title = 'Cálculo Individual'
        wb.properties.creator = self.author
        wb.properties.created = datetime.now()
        
        wb.save(filename)
        
        return filename
    
    def __str__(self) -> str:
        """Representación en string del modelo."""
        return f"ExcelExportModel(title='{self.title}', include_charts={self.include_charts})"
    
    def __repr__(self) -> str:
        """Representación detallada del modelo."""
        return f"ExcelExportModel(title={self.title!r}, author={self.author!r}, include_charts={self.include_charts})"

# Made with Bob
